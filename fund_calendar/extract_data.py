"""자금 현황 보고 엑셀(xlsx)에서 자금 캘린더용 데이터(JSON)를 뽑는다.

매출_YYMM / 지출_YYMM / 자금현황_YYMM 시트의 값만 원본 배열 형태로 옮기고,
은행·계좌번호·예금주처럼 캘린더에 필요 없는 민감 컬럼은 버린다.
시트 해석(확정/추정 판정 포함)은 index.html 쪽 파서가 담당한다.

사용법:
    python fund_calendar/extract_data.py 자금현황.xlsx -o fund_calendar/data/fund-data.json
    python fund_calendar/extract_data.py 자금현황.xlsx --embed fund_calendar/data/fund-calendar.html
"""
import argparse
import datetime as dt
import json
import re
from pathlib import Path

import openpyxl

SHEET_RE = re.compile(r"^(매출|지출|자금현황)_(\d{4}|\d{6})$")
DROP_HEADERS = {"은행", "계좌번호", "예금주"}
BALANCE_ROWS = 60  # 자금현황 시트는 상단 '1. 자금 종합' 표만 필요하다.
PLACEHOLDER = "/*__FUND_DATA__*/null"


def cell(v):
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def dump_sheet(ws, kind):
    max_row = BALANCE_ROWS if kind == "자금현황" else None
    rows = [[cell(v) for v in r] for r in ws.iter_rows(max_row=max_row, values_only=True)]
    if not rows:
        return []
    if kind != "자금현황":
        keep = [i for i, h in enumerate(rows[0]) if h not in DROP_HEADERS]
        rows = [[r[i] if i < len(r) else None for i in keep] for r in rows]
    out = []
    for r in rows:
        while r and r[-1] is None:
            r.pop()
        out.append(r)
    # 빈 행은 header 행 이후에만 제거한다(파서는 내용으로 행을 찾는다).
    return [out[0]] + [r for r in out[1:] if r]


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = {}
    for name in wb.sheetnames:
        m = SHEET_RE.match(name)
        if m:
            sheets[name] = dump_sheet(wb[name], m.group(1))
    return {
        "source": Path(path).name,
        "generatedAt": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "sheets": sheets,
    }


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("-o", "--output", help="JSON 파일 경로")
    ap.add_argument("--embed", help="데이터를 내장한 단일 HTML 파일 경로")
    ap.add_argument("--sheet-url", help="'원본 시트 열기' 링크로 쓸 구글 시트 URL")
    args = ap.parse_args()

    data = extract(args.xlsx)
    if args.sheet_url:
        data["sheetUrl"] = args.sheet_url
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":"))

    if args.output:
        Path(args.output).parent.mkdir(parents=True, exist_ok=True)
        Path(args.output).write_text(payload, encoding="utf-8")
    if args.embed:
        template = (Path(__file__).parent / "index.html").read_text(encoding="utf-8")
        if PLACEHOLDER not in template:
            raise SystemExit("index.html에서 데이터 자리표시자를 찾지 못했습니다.")
        safe = payload.replace("</", "<\\/")
        Path(args.embed).parent.mkdir(parents=True, exist_ok=True)
        Path(args.embed).write_text(template.replace(PLACEHOLDER, safe), encoding="utf-8")
    if not (args.output or args.embed):
        print(payload)


if __name__ == "__main__":
    main()
